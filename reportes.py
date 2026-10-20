"""Reportes consolidados por día, semana o mes, y exportación a Excel."""
import io
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from activos_calc import activos_vigentes, depreciacion_dia
from models import (Ajuste, Activo, CompraInsumo, Gasto, Produccion, Producto, Venta, VentaItem, db)

MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
AGRUPACIONES = {"diario": "Por día", "semanal": "Por semana", "mensual": "Por mes"}

COLUMNAS = [  # clave, título, es dinero
    ("compras", "Compras insumos", True),
    ("compras_eq", "Compras equipos", True),
    ("unid_prod", "Vikingos producidos", False),
    ("costo_prod", "Costo producción", True),
    ("unid_vend", "Vikingos vendidos", False),
    ("num_ventas", "N.º ventas", False),
    ("ventas", "Ventas", True),
    ("costo_vend", "Costo de lo vendido", True),
    ("bruta", "Ganancia bruta", True),
    ("equipos", "Depreciación equipos", True),
    ("gastos", "Gastos", True),
    ("neta", "Ganancia neta", True),
]


def inicio_bucket(d, agrupar):
    if agrupar == "semanal":
        return d - timedelta(days=d.weekday())
    if agrupar == "mensual":
        return d.replace(day=1)
    return d


def etiqueta(d, agrupar):
    if agrupar == "semanal":
        fin = d + timedelta(days=6)
        sem = d.isocalendar()[1]
        if d.month == fin.month:
            return f"Sem {sem} · {d.day}–{fin.day} {MESES[d.month - 1]}"
        return f"Sem {sem} · {d.day} {MESES[d.month - 1]}–{fin.day} {MESES[fin.month - 1]}"
    if agrupar == "mensual":
        return f"{MESES[d.month - 1].capitalize()} {d.year}"
    return d.strftime("%d/%m/%Y")


def rango_por_defecto(agrupar, hoy):
    if agrupar == "mensual":
        m, y = hoy.month - 5, hoy.year
        while m < 1:
            m += 12
            y -= 1
        return date(y, m, 1), hoy
    if agrupar == "semanal":
        return inicio_bucket(hoy, "semanal") - timedelta(weeks=7), hoy
    return hoy.replace(day=1), hoy


def generar(desde, hasta, agrupar):
    """desde/hasta: date (inclusive). Devuelve dict con filas, totales y detalles."""
    ini = datetime.combine(desde, datetime.min.time())
    fin = datetime.combine(hasta + timedelta(days=1), datetime.min.time())
    filas = defaultdict(lambda: defaultdict(float))

    def b(d):
        return inicio_bucket(d, agrupar)

    # Ventas
    for v in Venta.query.filter(Venta.anulada.is_(False), Venta.fecha >= ini, Venta.fecha < fin):
        f = filas[b(v.fecha.date())]
        f["ventas"] += v.total
        f["costo_vend"] += v.costo_total
        f["unid_vend"] += v.unidades
        f["num_ventas"] += 1
    # Compras
    for c in CompraInsumo.query.filter(CompraInsumo.anulada.is_(False), CompraInsumo.fecha >= ini,
                                       CompraInsumo.fecha < fin):
        filas[b(c.fecha.date())]["compras"] += c.costo_total
    for a in Activo.query.filter(Activo.eliminado.is_(False), Activo.fecha_compra >= desde,
                                 Activo.fecha_compra <= hasta):
        filas[b(a.fecha_compra)]["compras_eq"] += a.valor
    # Producción
    for p in Produccion.query.filter(Produccion.anulada.is_(False), Produccion.fecha >= ini,
                                     Produccion.fecha < fin):
        f = filas[b(p.fecha.date())]
        f["unid_prod"] += p.cantidad
        f["costo_prod"] += p.costo_total
    # Gastos
    gastos_cat = defaultdict(float)
    for g in Gasto.query.filter(Gasto.fecha >= desde, Gasto.fecha <= hasta):
        filas[b(g.fecha)]["gastos"] += g.valor
        gastos_cat[g.categoria] += g.valor
    # Egresos por ajuste (pérdidas de inventario): cuentan como gasto
    for a in Ajuste.query.filter(Ajuste.anulada.is_(False), Ajuste.fecha >= ini, Ajuste.fecha < fin):
        filas[b(a.fecha.date())]["gastos"] += a.total
        gastos_cat["Pérdidas de inventario"] += a.total
    # Depreciación (día a día)
    activos = activos_vigentes()
    d = desde
    while d <= hasta:
        dep = depreciacion_dia(d, activos)
        if dep:
            filas[b(d)]["equipos"] += dep
        d += timedelta(days=1)

    # Todos los periodos del rango, aunque estén vacíos
    claves, k = [], b(desde)
    while k <= hasta:
        claves.append(k)
        if agrupar == "mensual":
            k = (k.replace(day=28) + timedelta(days=4)).replace(day=1)
        elif agrupar == "semanal":
            k += timedelta(weeks=1)
        else:
            k += timedelta(days=1)

    salida, tot = [], defaultdict(float)
    for k in claves:
        f = filas[k]
        f["bruta"] = f["ventas"] - f["costo_vend"]
        f["neta"] = f["bruta"] - f["equipos"] - f["gastos"]
        fila = {c: f[c] for c, _, _ in COLUMNAS}
        fila["margen"] = (f["neta"] / f["ventas"] * 100) if f["ventas"] else None
        fila["periodo"] = etiqueta(k, agrupar)
        fila["inicio"] = k
        salida.append(fila)
        for c, _, _ in COLUMNAS:
            tot[c] += f[c]
    tot["margen"] = (tot["neta"] / tot["ventas"] * 100) if tot["ventas"] else None

    # Por referencia
    refs = (db.session.query(Producto.nombre, func.sum(VentaItem.cantidad),
                             func.sum(VentaItem.subtotal), func.sum(VentaItem.costo))
            .select_from(VentaItem).join(Producto, VentaItem.producto_id == Producto.id)
            .join(Venta, VentaItem.venta_id == Venta.id)
            .filter(Venta.anulada.is_(False), Venta.fecha >= ini, Venta.fecha < fin)
            .group_by(Producto.nombre).order_by(func.sum(VentaItem.cantidad).desc()).all())
    por_ref = [dict(nombre=r[0], unidades=int(r[1] or 0), ventas=r[2] or 0,
                    costo=r[3] or 0, ganancia=(r[2] or 0) - (r[3] or 0)) for r in refs]
    clis = (db.session.query(func.coalesce(Venta.cliente, "Cliente general"), func.count(func.distinct(Venta.id)),
                             func.sum(VentaItem.cantidad), func.sum(VentaItem.subtotal))
            .select_from(Venta).join(VentaItem, VentaItem.venta_id == Venta.id)
            .filter(Venta.anulada.is_(False), Venta.fecha >= ini, Venta.fecha < fin)
            .group_by(func.coalesce(Venta.cliente, "Cliente general"))
            .order_by(func.sum(VentaItem.cantidad).desc()).all())
    por_cliente = [dict(nombre=r[0], compras=int(r[1] or 0), unidades=int(r[2] or 0), ventas=r[3] or 0)
                   for r in clis]
    return dict(filas=salida, tot=dict(tot), por_ref=por_ref, por_cliente=por_cliente,
                gastos_cat=sorted(gastos_cat.items(), key=lambda x: -x[1]))


# ------------------------------------------------------------------ Excel
AZUL = "0F2A44"
CELESTE = "E2F6FA"
FMT_COP = '"$" #,##0;[Red]-"$" #,##0'
FMT_NUM = '#,##0'
FMT_PCT = '0.0"%"'


def _encabezado(ws, fila, titulos):
    borde = Border(bottom=Side(style="medium", color=AZUL))
    for i, t in enumerate(titulos, 1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borde
    ws.row_dimensions[fila].height = 32


def _anchos(ws, anchos):
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def excel(rep, desde, hasta, agrupar, negocio):
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"
    ws["A1"] = f"{negocio} — Reporte {AGRUPACIONES[agrupar].lower()}"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws["A2"] = f"Del {desde.strftime('%d/%m/%Y')} al {hasta.strftime('%d/%m/%Y')}"
    ws["A2"].font = Font(color="5F6F84")

    titulos = ["Periodo"] + [t for _, t, _ in COLUMNAS] + ["Margen neto"]
    _encabezado(ws, 4, titulos)
    r = 5
    for f in rep["filas"]:
        ws.cell(row=r, column=1, value=f["periodo"])
        for j, (c, _, dinero) in enumerate(COLUMNAS, 2):
            cell = ws.cell(row=r, column=j, value=round(f[c], 2))
            cell.number_format = FMT_COP if dinero else FMT_NUM
        m = ws.cell(row=r, column=len(COLUMNAS) + 2, value=round(f["margen"], 1) if f["margen"] is not None else None)
        m.number_format = FMT_PCT
        if r % 2 == 0:
            for j in range(1, len(titulos) + 1):
                ws.cell(row=r, column=j).fill = PatternFill("solid", fgColor="F5F8FB")
        r += 1
    ultima = r - 1
    # Totales con fórmulas
    ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    for j, (c, _, dinero) in enumerate(COLUMNAS, 2):
        col = get_column_letter(j)
        cell = ws.cell(row=r, column=j, value=f"=SUM({col}5:{col}{ultima})")
        cell.number_format = FMT_COP if dinero else FMT_NUM
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=CELESTE)
    col_v = get_column_letter(1 + [c for c, _, _ in COLUMNAS].index("ventas") + 1)
    col_n = get_column_letter(1 + [c for c, _, _ in COLUMNAS].index("neta") + 1)
    m = ws.cell(row=r, column=len(COLUMNAS) + 2, value=f'=IF({col_v}{r}=0,"",{col_n}{r}/{col_v}{r}*100)')
    m.number_format = FMT_PCT
    m.font = Font(bold=True)
    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=CELESTE)
    ws.cell(row=r, column=len(COLUMNAS) + 2).fill = PatternFill("solid", fgColor=CELESTE)
    ws.freeze_panes = "B5"
    _anchos(ws, [24] + [15] * len(COLUMNAS) + [12])

    # Gráfico ventas vs ganancia neta
    if len(rep["filas"]) > 1:
        idx_v = [c for c, _, _ in COLUMNAS].index("ventas") + 2
        idx_n = [c for c, _, _ in COLUMNAS].index("neta") + 2
        ch = BarChart()
        ch.title = "Ventas y ganancia neta"
        ch.y_axis.numFmt = '"$" #,##0'
        ch.add_data(Reference(ws, min_col=idx_v, min_row=4, max_row=ultima), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=ultima))
        ln = LineChart()
        ln.add_data(Reference(ws, min_col=idx_n, min_row=4, max_row=ultima), titles_from_data=True)
        ch += ln
        ch.height, ch.width = 9, 22
        ws.add_chart(ch, f"B{r + 3}")

    # Por referencia
    w2 = wb.create_sheet("Por referencia")
    _encabezado(w2, 1, ["Referencia", "Unidades vendidas", "Ventas", "Costo", "Ganancia bruta", "Margen"])
    for i, p in enumerate(rep["por_ref"], 2):
        w2.cell(row=i, column=1, value=p["nombre"])
        w2.cell(row=i, column=2, value=p["unidades"]).number_format = FMT_NUM
        for j, k in ((3, "ventas"), (4, "costo"), (5, "ganancia")):
            w2.cell(row=i, column=j, value=round(p[k], 2)).number_format = FMT_COP
        w2.cell(row=i, column=6, value=f'=IF(C{i}=0,"",E{i}/C{i}*100)').number_format = FMT_PCT
    _anchos(w2, [32, 16, 16, 16, 16, 10])

    # Por cliente
    wc = wb.create_sheet("Por cliente")
    _encabezado(wc, 1, ["Cliente", "Compras", "Vikingos", "Ventas"])
    for i, cl in enumerate(rep.get("por_cliente", []), 2):
        wc.cell(row=i, column=1, value=cl["nombre"])
        wc.cell(row=i, column=2, value=cl["compras"]).number_format = FMT_NUM
        wc.cell(row=i, column=3, value=cl["unidades"]).number_format = FMT_NUM
        wc.cell(row=i, column=4, value=round(cl["ventas"], 2)).number_format = FMT_COP
    _anchos(wc, [32, 12, 12, 16])
    wc.freeze_panes = "A2"

    # Gastos por categoría
    w3 = wb.create_sheet("Gastos")
    _encabezado(w3, 1, ["Categoría", "Total"])
    for i, (cat, v) in enumerate(rep["gastos_cat"], 2):
        w3.cell(row=i, column=1, value=cat)
        w3.cell(row=i, column=2, value=round(v, 2)).number_format = FMT_COP
    detalle = (Gasto.query.filter(Gasto.fecha >= desde, Gasto.fecha <= hasta)
               .order_by(Gasto.fecha).all())
    base = len(rep["gastos_cat"]) + 4
    _encabezado(w3, base, ["Fecha", "Categoría", "Descripción", "Valor"])
    for i, g in enumerate(detalle, base + 1):
        w3.cell(row=i, column=1, value=g.fecha).number_format = "DD/MM/YYYY"
        w3.cell(row=i, column=2, value=g.categoria)
        w3.cell(row=i, column=3, value=g.descripcion or "")
        w3.cell(row=i, column=4, value=round(g.valor, 2)).number_format = FMT_COP
    fila_aj = w3.max_row + 1
    ini_aj = datetime.combine(desde, datetime.min.time())
    fin_aj = datetime.combine(hasta + timedelta(days=1), datetime.min.time())
    for a in (Ajuste.query.filter(Ajuste.anulada.is_(False), Ajuste.fecha >= ini_aj, Ajuste.fecha < fin_aj)
              .order_by(Ajuste.fecha).all()):
        w3.cell(row=fila_aj, column=1, value=a.fecha.date()).number_format = "DD/MM/YYYY"
        w3.cell(row=fila_aj, column=2, value="Pérdidas de inventario")
        det = ", ".join(f"{it.cantidad:g} {it.unidad} {it.nombre}" for it in a.items)
        w3.cell(row=fila_aj, column=3, value=f"EG-{a.id:06d} · {a.motivo}: {det}"[:250])
        w3.cell(row=fila_aj, column=4, value=round(a.total, 2)).number_format = FMT_COP
        fila_aj += 1
    _anchos(w3, [24, 22, 40, 16])

    # Compras detalladas
    w4 = wb.create_sheet("Compras")
    _encabezado(w4, 1, ["Fecha", "Insumo", "Cantidad", "Unidad", "Costo unitario", "Total", "Proveedor"])
    ini = datetime.combine(desde, datetime.min.time())
    fin = datetime.combine(hasta + timedelta(days=1), datetime.min.time())
    for i, c in enumerate(CompraInsumo.query.filter(CompraInsumo.anulada.is_(False),
                                                    CompraInsumo.fecha >= ini, CompraInsumo.fecha < fin)
                          .order_by(CompraInsumo.fecha).all(), 2):
        w4.cell(row=i, column=1, value=c.fecha.date()).number_format = "DD/MM/YYYY"
        w4.cell(row=i, column=2, value=c.insumo.nombre)
        w4.cell(row=i, column=3, value=c.cantidad).number_format = "#,##0.###"
        w4.cell(row=i, column=4, value=c.insumo.unidad)
        w4.cell(row=i, column=5, value=round(c.costo_unitario, 2)).number_format = FMT_COP
        w4.cell(row=i, column=6, value=round(c.costo_total, 2)).number_format = FMT_COP
        w4.cell(row=i, column=7, value=c.proveedor or "")
    fila_eq = w4.max_row + 1
    for a in (Activo.query.filter(Activo.eliminado.is_(False), Activo.fecha_compra >= desde,
                                  Activo.fecha_compra <= hasta).order_by(Activo.fecha_compra).all()):
        w4.cell(row=fila_eq, column=1, value=a.fecha_compra).number_format = "DD/MM/YYYY"
        w4.cell(row=fila_eq, column=2, value=f"Equipo: {a.nombre}")
        w4.cell(row=fila_eq, column=3, value=1)
        w4.cell(row=fila_eq, column=4, value="equipo")
        w4.cell(row=fila_eq, column=5, value=round(a.valor, 2)).number_format = FMT_COP
        w4.cell(row=fila_eq, column=6, value=round(a.valor, 2)).number_format = FMT_COP
        w4.cell(row=fila_eq, column=7, value=a.proveedor or "")
        fila_eq += 1
    _anchos(w4, [12, 32, 12, 12, 15, 15, 26])

    # Ventas detalladas
    w5 = wb.create_sheet("Ventas")
    _encabezado(w5, 1, ["Fecha", "Comprobante", "Cliente", "Pago", "Referencia", "Cantidad",
                        "Precio", "Subtotal", "Costo", "Ganancia"])
    i = 2
    for v in (Venta.query.filter(Venta.anulada.is_(False), Venta.fecha >= ini, Venta.fecha < fin)
              .order_by(Venta.fecha).all()):
        for it in v.items:
            w5.cell(row=i, column=1, value=v.fecha.date()).number_format = "DD/MM/YYYY"
            w5.cell(row=i, column=2, value=f"VK-{v.id:06d}")
            w5.cell(row=i, column=3, value=v.cliente or "")
            w5.cell(row=i, column=4, value=v.metodo_pago or "")
            w5.cell(row=i, column=5, value=it.producto.nombre)
            w5.cell(row=i, column=6, value=it.cantidad)
            for j, val in ((7, it.precio_unitario), (8, it.subtotal), (9, it.costo),
                           (10, it.subtotal - it.costo)):
                w5.cell(row=i, column=j, value=round(val, 2)).number_format = FMT_COP
            i += 1
    _anchos(w5, [12, 14, 22, 14, 28, 10, 12, 14, 14, 14])
    for w in (w2, w3, w4, w5):
        w.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
